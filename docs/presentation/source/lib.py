"""Помощники для заполнения шаблона ФСП: текст с сохранением стиля, картинки, порядок слайдов."""

import copy
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

PURPLE = RGBColor(0x52, 0x09, 0x78)
DEEP = RGBColor(0x31, 0x0F, 0x53)
PINK = RGBColor(0xFF, 0x00, 0x53)
ROSE = RGBColor(0xFC, 0x37, 0x77)
LILAC = RGBColor(0x8A, 0x83, 0xD1)
BLUSH = RGBColor(0xFF, 0xD6, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x1D, 0x22)
GRAY = RGBColor(0x6B, 0x6C, 0x75)


def shapes_by_name(slide, name):
    return [s for s in slide.shapes if s.name == name]


def at(slide, x, y, tol=0.08):
    """Фигура по координатам левого верхнего угла в дюймах."""
    found = [
        s
        for s in slide.shapes
        if s.left is not None and abs(Emu(s.left).inches - x) < tol and abs(Emu(s.top).inches - y) < tol
    ]
    if not found:
        raise KeyError(f'нет фигуры в ({x}, {y})')
    return found[0]


def remove(shape):
    el = shape._element
    el.getparent().remove(el)


def _runs(text):
    """«**жирный** обычный» → [(текст, жирный)]."""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    out = []
    for p in parts:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            out.append((p[2:-2], True))
        else:
            out.append((p, False))
    return out


def fill(shape, paras, size=None, color=None, bold=None, align=None, autofit=False, tpl=None):
    """Заменить текст, сохранив оформление абзацев шаблона.

    paras — строка или список; элемент списка — строка (с **жирным**) или dict
    {t, size, color, bold, tpl} где tpl — номер абзаца шаблона для оформления.
    """
    if isinstance(paras, str):
        paras = [paras]
    body = shape.text_frame._txBody
    old = body.findall(qn('a:p'))
    for p in old:
        body.remove(p)
    for i, spec in enumerate(paras):
        if isinstance(spec, str):
            spec = {'t': spec}
        src = old[min(spec.get('tpl', i if tpl is None else tpl), len(old) - 1)]
        p = copy.deepcopy(src)
        base = src.find(qn('a:r') + '/' + qn('a:rPr'))
        if base is None:
            base = src.find(qn('a:endParaRPr'))
        for child in list(p):
            if child.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
                p.remove(child)
        end = p.find(qn('a:endParaRPr'))
        for text, is_bold in _runs(spec['t']):
            r = etree.SubElement(p, qn('a:r'))
            rpr = copy.deepcopy(base) if base is not None else etree.Element(qn('a:rPr'))
            rpr.tag = qn('a:rPr')
            rpr.set('lang', 'ru-RU')
            b = spec.get('bold', bold)
            if is_bold or b:
                rpr.set('b', '1')
            elif 'b' in rpr.attrib:
                del rpr.attrib['b']
            sz = spec.get('size', size)
            if sz:
                rpr.set('sz', str(int(sz * 100)))
            col = spec.get('color', color)
            if col is not None:
                for f in rpr.findall(qn('a:solidFill')):
                    rpr.remove(f)
                sf = etree.Element(qn('a:solidFill'))
                c = etree.SubElement(sf, qn('a:srgbClr'))
                c.set('val', str(col))
                rpr.insert(0, sf)
            r.append(rpr)
            t = etree.SubElement(r, qn('a:t'))
            t.text = text
            if end is not None:
                end.addprevious(r)
        if spec.get('space'):
            ppr = p.find(qn('a:pPr'))
            if ppr is None:
                ppr = etree.Element(qn('a:pPr'))
                p.insert(0, ppr)
            for c in ppr.findall(qn('a:spcBef')):
                ppr.remove(c)
            sb = etree.Element(qn('a:spcBef'))
            etree.SubElement(sb, qn('a:spcPts')).set('val', str(int(spec['space'] * 100)))
            ppr.insert(0, sb)
        if align is not None:
            ppr = p.find(qn('a:pPr'))
            if ppr is None:
                ppr = etree.SubElement(p, qn('a:pPr'))
                p.insert(0, ppr)
            ppr.set('algn', {'l': 'l', 'c': 'ctr', 'r': 'r'}[align])
        body.append(p)
    if autofit:
        bp = body.find(qn('a:bodyPr'))
        for child in list(bp):
            if child.tag in (qn('a:spAutoFit'), qn('a:noAutofit'), qn('a:normAutofit')):
                bp.remove(child)
        etree.SubElement(bp, qn('a:normAutofit'))


def textbox(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align='l', anchor='t', font=None, spacing=1.1):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        if isinstance(spec, str):
            spec = {'t': spec}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[spec.get('align', align)]
        p.line_spacing = spacing
        if spec.get('space'):
            p.space_before = Pt(spec['space'])
        for text, is_bold in _runs(spec['t']):
            r = p.add_run()
            r.text = text
            r.font.size = Pt(spec.get('size', size))
            r.font.bold = is_bold or spec.get('bold', bold)
            r.font.color.rgb = spec.get('color', color)
            if font or spec.get('font'):
                r.font.name = spec.get('font', font)
    return tb


def box(slide, x, y, w, h, fill_color=None, line=None, radius=0.12, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill_color is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill_color
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def reorder(prs, order):
    """Оставить слайды шаблона в порядке order (номера с 1), остальные удалить."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep = [ids[n - 1] for n in order]
    for el in ids:
        lst.remove(el)
        if el not in keep:
            prs.part.drop_rel(el.get(qn('r:id')))
    for el in keep:
        lst.append(el)
